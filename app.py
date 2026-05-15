import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt
import smtplib
from email.message import EmailMessage
import tempfile

st.set_page_config(page_title="Sellout Report")

st.title("📊 Sell-In / Sell-Out Report Generator")

uploaded_file = st.file_uploader(
    "Carica file CSV o Excel",
    type=["csv", "xlsx"]
)

email_dest = st.text_input(
    "Inserisci email destinatario"
)

GMAIL_USER = "selloutreportsvgs@gmail.com"
GMAIL_PASSWORD = "wndlkafpxndjvbwu"


def generate_report(df):

    # Pulizia colonne
    df.columns = df.columns.astype(str)
    df.columns = df.columns.str.strip()

    # Debug colonne
    st.write("Colonne trovate nel file:")
    st.write(list(df.columns))

    vendite_col = None
    carichi_col = None
    valore_col = None

    brand_col = df.columns[0]

    # Ricerca automatica colonne
    for col in df.columns:

        nome = col.lower()

        if (
            "vend" in nome and
            ("q" in nome or "qt" in nome)
        ):
            vendite_col = col

        if (
            "caric" in nome and
            ("q" in nome or "qt" in nome)
        ):
            carichi_col = col

        if (
            "val" in nome and
            "vend" in nome
        ):
            valore_col = col

    st.write("Colonna Brand:", brand_col)
    st.write("Colonna Vendite:", vendite_col)
    st.write("Colonna Carichi:", carichi_col)
    st.write("Colonna Valore:", valore_col)

    # Controllo sicurezza
    if vendite_col is None:
        st.error("❌ Colonna vendite non trovata")
        st.stop()

    if carichi_col is None:
        st.error("❌ Colonna carichi non trovata")
        st.stop()

    # Conversione numerica
    for col in [vendite_col, carichi_col]:

        df[col] = (
            df[col]
            .astype(str)
            .str.replace(".", "", regex=False)
            .str.replace(",", ".", regex=False)
        )

        df[col] = pd.to_numeric(
            df[col],
            errors="coerce"
        )

    # Sell Through
    df["Sell Through %"] = (
        df[vendite_col] /
        df[carichi_col] * 100
    ).round(1)

    # Top performance
    top = df.sort_values(
        "Sell Through %",
        ascending=False
    ).head(10)

    # Grafico
    plt.figure(figsize=(8, 4))

    plt.bar(
        top[brand_col],
        top["Sell Through %"]
    )

    plt.xticks(rotation=45)

    plt.tight_layout()

    chart_path = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".png"
    ).name

    plt.savefig(chart_path)

    plt.close()

    # Documento Word
    doc = Document()

    doc.add_heading(
        "Report Sell-In / Sell-Out",
        level=1
    )

    media_sell = df["Sell Through %"].mean()

    doc.add_paragraph(
        f"Sell-through medio: {media_sell:.1f}%"
    )

    table = doc.add_table(
        rows=1,
        cols=3
    )

    hdr = table.rows[0].cells

    hdr[0].text = "Brand"
    hdr[1].text = "Sell Through"
    hdr[2].text = "Vendite"

    for _, row in top.iterrows():

        cells = table.add_row().cells

        cells[0].text = str(
            row[brand_col]
        )

        cells[1].text = (
            f"{row['Sell Through %']}%"
        )

        if valore_col:

            cells[2].text = str(
                row[valore_col]
            )

    doc.add_picture(
        chart_path,
        width=Inches(5)
    )

    report_path = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".docx"
    ).name

    doc.save(report_path)

    return report_path


if st.button("Genera Report"):

    if uploaded_file and email_dest:

        # Lettura file
        if uploaded_file.name.endswith(".csv"):

            df = pd.read_csv(
                uploaded_file,
                sep=";"
            )

        else:

            df = pd.read_excel(
                uploaded_file
            )

        # Genera report
        report = generate_report(df)

        # Invio mail
        msg = EmailMessage()

        msg["Subject"] = (
            "Report Sell-Out"
        )

        msg["From"] = GMAIL_USER

        msg["To"] = email_dest

        msg.set_content(
            "In allegato il report automatico."
        )

        with open(report, "rb") as f:

            msg.add_attachment(
                f.read(),
                maintype="application",
                subtype="octet-stream",
                filename="report.docx"
            )

        with smtplib.SMTP_SSL(
            "smtp.gmail.com",
            465
        ) as smtp:

            smtp.login(
                GMAIL_USER,
                GMAIL_PASSWORD
            )

            smtp.send_message(msg)

        st.success(
            "✅ Report inviato correttamente!"
        )
